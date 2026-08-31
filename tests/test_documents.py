from pathlib import Path
from zipfile import ZipFile

from docx import Document
from openpyxl import Workbook
from pptx import Presentation

from evolving_agent.commands import CommandRunner
from evolving_agent.documents import extract_text
from evolving_agent.tools import WorkspaceTools
from evolving_agent.workspace import Workspace


def test_extracts_docx_xlsx_and_pptx(tmp_path: Path) -> None:
    word_path = tmp_path / "brief.docx"
    document = Document()
    document.add_paragraph("Decision: approve the launch")
    document.save(word_path)
    assert "approve the launch" in extract_text(word_path)

    spreadsheet_path = tmp_path / "budget.xlsx"
    workbook = Workbook()
    workbook.active.title = "Costs"
    workbook.active.append(["Item", "Amount"])
    workbook.active.append(["Hosting", 42])
    workbook.save(spreadsheet_path)
    assert "[Sheet: Costs]" in extract_text(spreadsheet_path)
    assert "Hosting | 42" in extract_text(spreadsheet_path)

    slides_path = tmp_path / "update.pptx"
    slides = Presentation()
    slides.slides.add_slide(slides.slide_layouts[6]).shapes.add_textbox(0, 0, 1000000, 1000000).text_frame.text = "Milestone complete"
    slides.save(slides_path)
    assert "[Slide 1]" in extract_text(slides_path)
    assert "Milestone complete" in extract_text(slides_path)


def test_extracts_opendocument_text_and_spreadsheet(tmp_path: Path) -> None:
    odt = tmp_path / "brief.odt"
    with ZipFile(odt, "w") as archive:
        archive.writestr("content.xml", """<office:document-content xmlns:office='urn:o' xmlns:text='urn:t'>
        <office:body><office:text><text:h>Decision</text:h><text:p>Ship on Tuesday.</text:p></office:text></office:body>
        </office:document-content>""")
    assert "Decision" in extract_text(odt)
    assert "Ship on Tuesday." in extract_text(odt)

    ods = tmp_path / "budget.ods"
    with ZipFile(ods, "w") as archive:
        archive.writestr("content.xml", """<office:document-content xmlns:office='urn:o' xmlns:table='urn:table' xmlns:text='urn:t'>
        <office:body><office:spreadsheet><table:table table:name='Forecast'><table:table-row>
        <table:table-cell><text:p>Revenue</text:p></table:table-cell><table:table-cell><text:p>100</text:p></table:table-cell>
        </table:table-row></table:table></office:spreadsheet></office:body></office:document-content>""")
    result = extract_text(ods)
    assert "[Sheet: Forecast]" in result
    assert "Revenue | 100" in result


def test_read_document_can_read_materials_but_rejects_unknown_type(tmp_path: Path) -> None:
    root, materials = tmp_path / "work", tmp_path / "materials"
    root.mkdir()
    materials.mkdir()
    document = Document()
    document.add_paragraph("Confidential schedule")
    document.save(materials / "schedule.docx")
    tools = WorkspaceTools(Workspace(root), CommandRunner(root), materials=Workspace(materials))
    assert "Confidential schedule" in tools.call("read_document", {"path": "materials/schedule.docx"})
    (root / "notes.txt").write_text("not an office document")
    try:
        tools.call("read_document", {"path": "notes.txt"})
    except Exception as failure:
        assert "Unsupported document type" in str(failure)
    else:
        raise AssertionError("text files must not be parsed as office documents")
