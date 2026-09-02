"""Bounded, non-rendering inspection of PowerPoint presentation evidence.

The OOXML package is read through :mod:`python-pptx`; it is never rendered and
this module does not follow external relationships, open media, or execute any
active content.  It reports the text a reviewer can use, including speaker
notes and table cells, while putting fixed limits on input and output work.
"""

from __future__ import annotations

import json
import zipfile
from pathlib import Path
from typing import Any, Iterator

from pptx import Presentation

_MAX_SOURCE_BYTES = 32 * 1024 * 1024
_MAX_UNCOMPRESSED_BYTES = 128 * 1024 * 1024
_MAX_ARCHIVE_MEMBERS = 2_000
_MAX_SLIDES = 25
_MAX_TEXT_CHARACTERS = 12_000
_MAX_METADATA_VALUE = 1_000


class PresentationInspectionError(ValueError):
    """A presentation could not be safely inspected."""


def inspect_presentation(
    path: Path,
    *,
    slide: int | None = None,
    max_slides: int = 5,
    max_characters: int = 8_000,
) -> str:
    """Return core properties and bounded text from one or several PPTX slides.

    ``slide`` is one-based when provided. Speaker notes and table cells are
    represented as text, but embedded files, media, macros, and hyperlinks are
    not opened or followed.
    """
    if not 1 <= max_slides <= _MAX_SLIDES:
        raise PresentationInspectionError(f"max_slides must be between 1 and {_MAX_SLIDES}.")
    if not 1 <= max_characters <= _MAX_TEXT_CHARACTERS:
        raise PresentationInspectionError(
            f"max_characters must be between 1 and {_MAX_TEXT_CHARACTERS}."
        )
    if slide is not None and slide < 1:
        raise PresentationInspectionError("slide must be one or greater (slides are one-based).")
    if not path.is_file():
        raise PresentationInspectionError(f"{path.name!r} is not a readable PPTX file.")
    try:
        size = path.stat().st_size
    except OSError as exc:
        raise PresentationInspectionError(f"Could not inspect {path.name!r}: {exc}") from exc
    if size > _MAX_SOURCE_BYTES:
        raise PresentationInspectionError(
            f"{path.name!r} is {size} bytes; PPTX inspection is limited to {_MAX_SOURCE_BYTES} bytes."
        )
    try:
        # is_zipfile prevents python-pptx's less actionable package error for
        # ordinary non-PPTX input without extracting or executing anything.
        if not zipfile.is_zipfile(path):
            raise PresentationInspectionError(f"{path.name!r} is not a ZIP-based PPTX file.")
        with zipfile.ZipFile(path) as package:
            members = package.infolist()
            total_uncompressed = sum(member.file_size for member in members)
        if len(members) > _MAX_ARCHIVE_MEMBERS:
            raise PresentationInspectionError(
                f"{path.name!r} has too many package members (limit {_MAX_ARCHIVE_MEMBERS})."
            )
        if total_uncompressed > _MAX_UNCOMPRESSED_BYTES:
            raise PresentationInspectionError(
                f"{path.name!r} expands to {total_uncompressed} bytes; PPTX inspection is limited to "
                f"{_MAX_UNCOMPRESSED_BYTES} uncompressed bytes."
            )
        presentation = Presentation(path)
        slide_count = len(presentation.slides)
    except PresentationInspectionError:
        raise
    except (OSError, KeyError, ValueError, zipfile.BadZipFile) as exc:
        raise PresentationInspectionError(f"Could not parse PPTX {path.name!r}: {exc}") from exc

    if slide is not None and slide > slide_count:
        raise PresentationInspectionError(f"slide {slide} is outside this presentation ({slide_count} slides).")
    indices = range(slide - 1, slide) if slide is not None else range(min(slide_count, max_slides))
    report: dict[str, Any] = {
        "format": "pptx",
        "size_bytes": size,
        "slide_count": slide_count,
        "metadata": _metadata(presentation),
        "slides": [],
    }
    for index in indices:
        item = presentation.slides[index]
        slide_text = "\n".join(_shape_text(item.shapes))
        notes_text = _notes_text(item)
        entry: dict[str, Any] = {
            "slide_number": index + 1,
            "text_preview": slide_text[:max_characters],
            "text_preview_truncated": len(slide_text) > max_characters,
        }
        if notes_text:
            entry["notes_preview"] = notes_text[:max_characters]
            entry["notes_preview_truncated"] = len(notes_text) > max_characters
        report["slides"].append(entry)
    report["slides_truncated"] = slide is None and slide_count > len(report["slides"])
    report["note"] = (
        "Text is read from slide shapes, tables, groups, and speaker notes only; "
        "the presentation is not rendered and macros, media, embedded files, and links are not opened."
    )
    return json.dumps(report, ensure_ascii=False, indent=2)


def _shape_text(shapes: Any) -> Iterator[str]:
    """Yield non-empty text from a shape tree, including grouped table cells."""
    for shape in shapes:
        if getattr(shape, "has_table", False):
            for row in shape.table.rows:
                yield "\t".join(cell.text for cell in row.cells)
        elif getattr(shape, "has_text_frame", False):
            text = shape.text
            if text:
                yield text
        if getattr(shape, "shape_type", None) is not None and hasattr(shape, "shapes"):
            yield from _shape_text(shape.shapes)


def _notes_text(slide: Any) -> str:
    """Return speaker-note text, omitting empty placeholders when unavailable."""
    try:
        return "\n".join(_shape_text(slide.notes_slide.shapes))
    except (AttributeError, KeyError, ValueError):
        return ""


def _metadata(presentation: Any) -> dict[str, str]:
    properties = presentation.core_properties
    result: dict[str, str] = {}
    for name in ("title", "subject", "author", "keywords", "comments", "category", "last_modified_by"):
        value = getattr(properties, name, None)
        if value:
            result[name] = _bounded(str(value))
    return result


def _bounded(value: str) -> str:
    value = " ".join(value.split())
    return value if len(value) <= _MAX_METADATA_VALUE else value[:_MAX_METADATA_VALUE] + "…"
