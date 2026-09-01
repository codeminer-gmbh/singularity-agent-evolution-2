"""Bounded creation of ordinary XLSX workbook deliverables.

The writer takes data already selected by the model and writes it atomically.  It
never opens an input workbook, preserves no macros, and does not calculate
formulas; formulas are stored for Excel or another spreadsheet application to
calculate when the recipient opens the file.
"""
from __future__ import annotations

import os
import tempfile
from collections.abc import Mapping, Sequence
from pathlib import Path
from typing import Any

from openpyxl import Workbook

_MAX_SHEETS = 25
_MAX_ROWS = 2_000
_MAX_COLUMNS = 100
_MAX_CELLS = 100_000
_MAX_TEXT = 500_000
_MAX_CELL_TEXT = 20_000
_MAX_RESULT_BYTES = 10 * 1024 * 1024
_BAD_SHEET_CHARS = set("[]:*?/\\\\")


class XlsxWriterError(ValueError):
    """The requested workbook is outside the deliberately small writer contract."""


def create_xlsx(path: Path, *, sheets: object) -> dict[str, int]:
    """Create an XLSX with named sheets, rectangular-ish rows, and basic view settings."""
    if path.suffix.lower() != ".xlsx":
        raise XlsxWriterError("The deliverable path must end in .xlsx.")
    if not isinstance(sheets, Sequence) or isinstance(sheets, (str, bytes)) or not sheets:
        raise XlsxWriterError("'sheets' must be a non-empty array of sheet objects.")
    if len(sheets) > _MAX_SHEETS:
        raise XlsxWriterError(f"At most {_MAX_SHEETS} sheets may be created.")

    workbook = Workbook()
    workbook.remove(workbook.active)
    total_cells = total_text = 0
    names: set[str] = set()
    for sheet_index, specification in enumerate(sheets, 1):
        if not isinstance(specification, Mapping):
            raise XlsxWriterError(f"sheets[{sheet_index}] must be an object.")
        name = _sheet_name(specification.get("name"), sheet_index)
        if name.casefold() in names:
            raise XlsxWriterError(f"Sheet name {name!r} is duplicated.")
        names.add(name.casefold())
        rows = specification.get("rows")
        if not isinstance(rows, Sequence) or isinstance(rows, (str, bytes)) or not rows:
            raise XlsxWriterError(f"sheets[{sheet_index}].rows must be a non-empty array.")
        if len(rows) > _MAX_ROWS:
            raise XlsxWriterError(f"sheets[{sheet_index}] exceeds {_MAX_ROWS} rows.")
        worksheet = workbook.create_sheet(name)
        for row_index, row in enumerate(rows, 1):
            if not isinstance(row, Sequence) or isinstance(row, (str, bytes)):
                raise XlsxWriterError(f"sheets[{sheet_index}].rows[{row_index}] must be an array.")
            if len(row) > _MAX_COLUMNS:
                raise XlsxWriterError(f"Sheets may have at most {_MAX_COLUMNS} columns.")
            total_cells += len(row)
            if total_cells > _MAX_CELLS:
                raise XlsxWriterError(f"Workbook may contain at most {_MAX_CELLS} supplied cells.")
            for column_index, value in enumerate(row, 1):
                checked, characters = _cell_value(value, f"sheets[{sheet_index}].rows[{row_index}][{column_index}]")
                total_text += characters
                if total_text > _MAX_TEXT:
                    raise XlsxWriterError(f"Workbook text exceeds {_MAX_TEXT} characters.")
                worksheet.cell(row=row_index, column=column_index, value=checked)
        _configure_sheet(worksheet, specification, sheet_index)

    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        descriptor, temporary_name = tempfile.mkstemp(prefix=".xlsx-write-", suffix=".xlsx", dir=path.parent)
        os.close(descriptor)
        temporary = Path(temporary_name)
        try:
            workbook.save(temporary)
            size = temporary.stat().st_size
            if size > _MAX_RESULT_BYTES:
                raise XlsxWriterError("The generated XLSX exceeds the 10 MB deliverable limit.")
            os.replace(temporary, path)
        finally:
            temporary.unlink(missing_ok=True)
    except XlsxWriterError:
        raise
    except (OSError, ValueError) as exc:
        raise XlsxWriterError(f"Could not write {path.name!r}: {exc}") from exc
    return {"size_bytes": size, "sheets": len(sheets), "cells": total_cells}


def _sheet_name(value: object, index: int) -> str:
    if not isinstance(value, str) or not value.strip() or len(value) > 31 or any(c in _BAD_SHEET_CHARS for c in value):
        raise XlsxWriterError(f"sheets[{index}].name must be a non-blank Excel sheet name of at most 31 characters.")
    return value


def _cell_value(value: object, field: str) -> tuple[Any, int]:
    if value is None or isinstance(value, (bool, int, float)):
        return value, 0
    if not isinstance(value, str):
        raise XlsxWriterError(f"{field} must be a string, number, boolean, or null.")
    if len(value) > _MAX_CELL_TEXT:
        raise XlsxWriterError(f"{field} exceeds the {_MAX_CELL_TEXT}-character cell limit.")
    return value, len(value)


def _configure_sheet(worksheet: Any, specification: Mapping[str, object], index: int) -> None:
    permitted = {"name", "rows", "freeze_panes", "column_widths"}
    unknown = set(specification) - permitted
    if unknown:
        raise XlsxWriterError(f"sheets[{index}] has unsupported fields: {', '.join(sorted(unknown))}.")
    frozen = specification.get("freeze_panes")
    if frozen is not None:
        if not isinstance(frozen, str) or not frozen or len(frozen) > 12:
            raise XlsxWriterError(f"sheets[{index}].freeze_panes must be a cell reference such as 'A2'.")
        worksheet.freeze_panes = frozen
    widths = specification.get("column_widths")
    if widths is not None:
        if not isinstance(widths, Mapping) or len(widths) > _MAX_COLUMNS:
            raise XlsxWriterError(f"sheets[{index}].column_widths must map up to {_MAX_COLUMNS} column letters to widths.")
        for column, width in widths.items():
            if not isinstance(column, str) or not column.isalpha() or len(column) > 3 or not isinstance(width, (int, float)) or isinstance(width, bool) or not 1 <= width <= 100:
                raise XlsxWriterError(f"sheets[{index}].column_widths entries need column letters and widths from 1 to 100.")
            worksheet.column_dimensions[column.upper()].width = width
