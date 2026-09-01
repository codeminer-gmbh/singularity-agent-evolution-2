"""Bounded structural inspection of Excel workbooks supplied with a task."""

from __future__ import annotations

from datetime import date, datetime, time
from decimal import Decimal
import json
from pathlib import Path
from typing import Any
import zipfile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

MAX_WORKBOOK_BYTES = 48 * 1024 * 1024
MAX_WORKBOOK_MEMBERS = 10_000
MAX_WORKBOOK_UNCOMPRESSED_BYTES = 256 * 1024 * 1024
MAX_WORKBOOK_SHEETS = 40
MAX_INSPECT_ROWS = 5_000
MAX_INSPECT_COLUMNS = 100
MAX_INSPECT_CELLS = 20_000
MAX_CELL_TEXT = 4_000
MAX_INSPECT_OUTPUT = 120_000


class SpreadsheetError(Exception):
    """A workbook cannot safely be inspected."""


def inspect_workbook(path: Path, *, max_rows: int = 200, max_columns: int = 30) -> str:
    """Return JSON describing workbook layout and a bounded, typed cell sample.

    Formula text is deliberately returned rather than a cached value, allowing a
    model to distinguish calculated fields from entered data.  This is not a
    spreadsheet calculation engine; Excel formulas are not evaluated by
    openpyxl.
    """
    if not path.is_file() or path.suffix.lower() not in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
        raise SpreadsheetError("inspect_workbook supports an existing .xlsx, .xlsm, .xltx, or .xltm file.")
    if path.stat().st_size > MAX_WORKBOOK_BYTES:
        raise SpreadsheetError(f"{path.name!r} exceeds the {MAX_WORKBOOK_BYTES}-byte workbook limit.")
    if isinstance(max_rows, bool) or not isinstance(max_rows, int) or not 1 <= max_rows <= MAX_INSPECT_ROWS:
        raise SpreadsheetError(f"max_rows must be an integer from 1 through {MAX_INSPECT_ROWS}.")
    if isinstance(max_columns, bool) or not isinstance(max_columns, int) or not 1 <= max_columns <= MAX_INSPECT_COLUMNS:
        raise SpreadsheetError(f"max_columns must be an integer from 1 through {MAX_INSPECT_COLUMNS}.")
    _check_zip(path)
    try:
        workbook = load_workbook(path, read_only=False, data_only=False, keep_links=False)
    except Exception as error:
        raise SpreadsheetError(f"Could not parse workbook {path.name!r}: {error}") from error
    try:
        if len(workbook.worksheets) > MAX_WORKBOOK_SHEETS:
            raise SpreadsheetError(f"Workbook has {len(workbook.worksheets)} sheets; limit is {MAX_WORKBOOK_SHEETS}.")
        remaining = MAX_INSPECT_CELLS
        sheets: list[dict[str, Any]] = []
        for sheet in workbook.worksheets:
            rows = min(sheet.max_row or 1, max_rows)
            columns = min(sheet.max_column or 1, max_columns)
            available_rows = min(rows, remaining // max(columns, 1))
            sampled_rows: list[dict[str, Any]] = []
            for row in sheet.iter_rows(min_row=1, max_row=available_rows, min_col=1, max_col=columns):
                cells = [_cell_summary(cell) for cell in row if cell.value is not None]
                if cells:
                    sampled_rows.append({"row": row[0].row, "cells": cells})
            remaining -= available_rows * columns
            tables = [
                {"name": table.name, "ref": table.ref, "display_name": table.displayName}
                for table in sheet.tables.values()
            ]
            sheets.append({
                "name": sheet.title,
                "state": sheet.sheet_state,
                "declared_dimensions": sheet.calculate_dimension(),
                "max_row": sheet.max_row,
                "max_column": sheet.max_column,
                "sampled_rows": sampled_rows,
                "sample_limits": {"rows": available_rows, "columns": columns},
                "merged_ranges": [str(area) for area in list(sheet.merged_cells.ranges)[:200]],
                "freeze_panes": str(sheet.freeze_panes) if sheet.freeze_panes else None,
                "auto_filter": sheet.auto_filter.ref,
                "tables": tables[:100],
            })
        result = json.dumps({"workbook": path.name, "sheets": sheets}, ensure_ascii=False, indent=2, default=_json_value)
        if len(result) > MAX_INSPECT_OUTPUT:
            return result[:MAX_INSPECT_OUTPUT] + f"\n... [truncated at {MAX_INSPECT_OUTPUT} characters]"
        return result
    finally:
        workbook.close()


def _check_zip(path: Path) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            infos = archive.infolist()
    except (OSError, ValueError, zipfile.BadZipFile) as error:
        raise SpreadsheetError(f"Could not open workbook ZIP package: {error}") from error
    if len(infos) > MAX_WORKBOOK_MEMBERS:
        raise SpreadsheetError(f"Workbook has more than {MAX_WORKBOOK_MEMBERS} package members.")
    expanded = sum(info.file_size for info in infos)
    if expanded > MAX_WORKBOOK_UNCOMPRESSED_BYTES:
        raise SpreadsheetError(f"Workbook expands to {expanded} bytes; limit is {MAX_WORKBOOK_UNCOMPRESSED_BYTES}.")


def _cell_summary(cell: Any) -> dict[str, Any]:
    value = cell.value
    text = _json_value(value)
    if isinstance(text, str) and len(text) > MAX_CELL_TEXT:
        text = text[:MAX_CELL_TEXT] + f"... [truncated at {MAX_CELL_TEXT} characters]"
    return {
        "coordinate": cell.coordinate,
        "column": get_column_letter(cell.column),
        "value": text,
        "formula": value if isinstance(value, str) and value.startswith("=") else None,
        "number_format": cell.number_format,
    }


def _json_value(value: Any) -> Any:
    if isinstance(value, (datetime, date, time, Decimal)):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)
